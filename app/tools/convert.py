"""Dönüştürme araçları."""
from __future__ import annotations

import io
import re
import warnings

import pymupdf

from .. import backends
from ..util import UserError, paper_size, parse_pages, save_pdf, stem
from . import image_bytes, tool

# ======================= PDF'E DÖNÜŞTÜR =======================


def _place_images(images, ctx, enhance=None) -> pymupdf.Document:
    from PIL import Image, ImageOps

    size = ctx.opt("page_size", "fit")
    orient = ctx.opt("orientation", "auto")
    margin = {"none": 0, "small": 20, "big": 40}.get(ctx.opt("margin", "none"), 0)
    out = pymupdf.open()
    for p in images:
        with Image.open(p) as im:
            im = ImageOps.exif_transpose(im)
            if enhance:
                im = enhance(im)
            buf, w, h = image_bytes(im)
        if size == "fit":
            W, H = w + 2 * margin, h + 2 * margin
        else:
            landscape = (w > h) if orient == "auto" else orient == "landscape"
            W, H = paper_size(size, landscape)
        page = out.new_page(width=W, height=H)
        box = pymupdf.Rect(margin, margin, W - margin, H - margin)
        page.insert_image(box, stream=buf, keep_proportion=True)
    return out


@tool("images_to_pdf", kinds=("image",))
def images_to_pdf(ctx):
    if ctx.opt_bool("merge", True):
        doc = _place_images([e.path for e in ctx.entries], ctx)
        name = stem(ctx.first.name) + ".pdf" if len(ctx.entries) == 1 else "gorseller.pdf"
        return save_pdf(doc, ctx.out(name))
    paths = []
    for e in ctx.entries:
        doc = _place_images([e.path], ctx)
        paths.append(save_pdf(doc, ctx.out(stem(e.name) + ".pdf")))
    ctx.result_name = "gorseller.zip"
    return paths


def _find_document_quad(cv2, np, img):
    """Fotoğraftaki belge kenarlarını bulur (4 köşe) ya da None döner."""
    h, w = img.shape[:2]
    scale = 800 / max(h, w)
    small = cv2.resize(img, (int(w * scale), int(h * scale))) if scale < 1 else img.copy()
    scale = min(scale, 1)
    gray = cv2.cvtColor(small, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray, (5, 5), 0)
    edges = cv2.Canny(gray, 50, 150)
    edges = cv2.dilate(edges, np.ones((3, 3), np.uint8), iterations=2)
    contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    area_min = 0.25 * small.shape[0] * small.shape[1]
    for c in sorted(contours, key=cv2.contourArea, reverse=True)[:8]:
        if cv2.contourArea(c) < area_min:
            break
        approx = cv2.approxPolyDP(c, 0.02 * cv2.arcLength(c, True), True)
        if len(approx) == 4 and cv2.isContourConvex(approx):
            return approx.reshape(4, 2).astype("float32") / scale
    return None


def _warp(cv2, np, img, quad):
    s = quad.sum(axis=1)
    d = np.diff(quad, axis=1).ravel()
    tl, br = quad[np.argmin(s)], quad[np.argmax(s)]
    tr, bl = quad[np.argmin(d)], quad[np.argmax(d)]
    src = np.array([tl, tr, br, bl], dtype="float32")
    wa = np.linalg.norm(br - bl)
    wb = np.linalg.norm(tr - tl)
    ha = np.linalg.norm(tr - br)
    hb = np.linalg.norm(tl - bl)
    W, H = int(max(wa, wb)), int(max(ha, hb))
    dst = np.array([[0, 0], [W - 1, 0], [W - 1, H - 1], [0, H - 1]], dtype="float32")
    M = cv2.getPerspectiveTransform(src, dst)
    return cv2.warpPerspective(img, M, (W, H))


def _scan_enhancer(mode: str, crop: bool):
    def enhance(im):
        import cv2
        import numpy as np
        from PIL import Image

        arr = cv2.cvtColor(np.array(im.convert("RGB")), cv2.COLOR_RGB2BGR)
        if crop:
            quad = _find_document_quad(cv2, np, arr)
            if quad is not None:
                arr = _warp(cv2, np, arr, quad)
        if mode == "bw":
            g = cv2.cvtColor(arr, cv2.COLOR_BGR2GRAY)
            g = cv2.adaptiveThreshold(g, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 15)
            return Image.fromarray(g)
        if mode == "gray":
            g = cv2.cvtColor(arr, cv2.COLOR_BGR2GRAY)
            g = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(g)
            return Image.fromarray(g)
        if mode == "auto":
            lab = cv2.cvtColor(arr, cv2.COLOR_BGR2LAB)
            l, a, b = cv2.split(lab)
            l = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(l)
            arr = cv2.cvtColor(cv2.merge((l, a, b)), cv2.COLOR_LAB2BGR)
        return Image.fromarray(cv2.cvtColor(arr, cv2.COLOR_BGR2RGB))
    return enhance


@tool("scan", kinds=("image",))
def scan(ctx):
    enhance = _scan_enhancer(ctx.opt("filter", "auto"), ctx.opt_bool("auto_crop", True))
    if "page_size" not in ctx.options:
        ctx.options["page_size"] = "A4"
    doc = _place_images([e.path for e in ctx.entries], ctx, enhance=enhance)
    if ctx.opt_bool("ocr", False):
        from .optimize import _ocr_page
        from ..fonts import get_font
        font = get_font("arial")
        for page in doc:
            _ocr_page(page, ctx.opt("language", "tur+eng"), 300, font)
    return save_pdf(doc, ctx.out("tarama.pdf"))


def _office(ctx, kinds):
    paths = []
    for e in ctx.entries:
        if e.kind not in kinds:
            raise UserError(f"'{e.name}' bu dönüşüm için uygun değil.")
        dst = ctx.out(stem(e.name) + ".pdf")
        backends.office_to_pdf(e.path, e.kind, dst)
        paths.append(dst)
    if len(paths) > 1 and ctx.opt_bool("merge", False):
        out = pymupdf.open()
        for p in paths:
            with pymupdf.open(p) as d:
                out.insert_pdf(d)
        return save_pdf(out, ctx.out("donusturulmus.pdf"))
    ctx.result_name = "donusturulmus.zip"
    return paths


@tool("word_to_pdf", kinds=("word",))
def word_to_pdf(ctx):
    return _office(ctx, ("word",))


@tool("excel_to_pdf", kinds=("excel",))
def excel_to_pdf(ctx):
    return _office(ctx, ("excel",))


@tool("ppt_to_pdf", kinds=("powerpoint",))
def ppt_to_pdf(ctx):
    return _office(ctx, ("powerpoint",))


@tool("office_to_pdf", kinds=("word", "excel", "powerpoint"))
def office_to_pdf(ctx):
    return _office(ctx, ("word", "excel", "powerpoint"))


@tool("html_to_pdf", kinds=("html",), min_files=0)
def html_to_pdf(ctx):
    url = (ctx.opt("url") or "").strip()
    hf = ctx.opt_bool("header_footer", False)
    if url:
        if not re.match(r"^https?://", url, re.I):
            url = "https://" + url
        host = re.sub(r"^https?://", "", url).split("/")[0] or "sayfa"
        dst = ctx.out(re.sub(r"[^\w.-]", "_", host) + ".pdf")
        return backends.html_to_pdf(url, dst, header_footer=hf)
    if not ctx.entries:
        raise UserError("Bir web adresi yaz ya da HTML dosyası ekle.")
    paths = []
    for e in ctx.entries:
        dst = ctx.out(stem(e.name) + ".pdf")
        paths.append(backends.html_to_pdf(e.path.as_uri(), dst, header_footer=hf))
    ctx.result_name = "html.zip"
    return paths


# ======================= PDF'TEN DÖNÜŞTÜR =======================


@tool("pdf_to_images")
def pdf_to_images(ctx):
    mode = ctx.opt("mode", "pages")
    fmt = ctx.opt("format", "jpg")
    dpi = int(ctx.opt_num("dpi", 150))
    quality = int(ctx.opt_num("quality", 90))
    paths = []
    for e in ctx.entries:
        doc = ctx.open_pdf(e)
        base = stem(e.name)
        if mode == "extract":
            seen = set()
            for pno in range(doc.page_count):
                for img in doc.get_page_images(pno, full=True):
                    xref = img[0]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    try:
                        info = doc.extract_image(xref)
                    except Exception:
                        continue
                    if not info or info["width"] < 32 or info["height"] < 32:
                        continue
                    ext = info["ext"]
                    data = info["image"]
                    if img[1]:  # maske (şeffaflık) varsa PNG olarak birleştir
                        try:
                            pix = pymupdf.Pixmap(doc, xref)
                            mask = pymupdf.Pixmap(doc, img[1])
                            pix = pymupdf.Pixmap(pix, mask)
                            data, ext = pix.tobytes("png"), "png"
                        except Exception:
                            pass
                    if ext not in ("jpg", "jpeg", "png"):
                        try:
                            pix = pymupdf.Pixmap(doc, xref)
                            if pix.n - pix.alpha >= 4:
                                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                            data, ext = pix.tobytes("png"), "png"
                        except Exception:
                            pass
                    p = ctx.out(f"{base}_s{pno + 1}_gorsel{len(seen)}.{ext}")
                    p.write_bytes(data)
                    paths.append(p)
        else:
            pages = parse_pages(ctx.opt("pages", "all"), doc.page_count)
            for i in pages:
                pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
                if fmt == "png":
                    p = ctx.out(f"{base}_sayfa_{i + 1}.png")
                    pix.save(p)
                else:
                    p = ctx.out(f"{base}_sayfa_{i + 1}.jpg")
                    p.write_bytes(pix.tobytes("jpg", jpg_quality=quality))
                paths.append(p)
    if not paths:
        raise UserError("PDF'te çıkarılacak görsel bulunamadı.")
    ctx.result_name = f"{stem(ctx.first.name)}_gorseller.zip"
    return paths


@tool("pdf_to_word")
def pdf_to_word(ctx):
    engine = ctx.opt("engine", "pdf2docx")
    paths = []
    for e in ctx.entries:
        src = ctx.pdf_path(e)
        dst = ctx.out(stem(e.name) + ".docx")
        if engine == "word" and backends.capabilities()["word"]:
            backends.pdf_to_word_msword(src, dst)
        else:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                from pdf2docx import Converter
            cv = Converter(str(src))
            try:
                cv.convert(str(dst), multi_processing=False)
            finally:
                cv.close()
        paths.append(dst)
    ctx.result_name = "word.zip"
    return paths


def _strip_text(page_doc: pymupdf.Document) -> None:
    for p in page_doc:
        for b in p.get_text("dict")["blocks"]:
            if b["type"] == 0:
                p.add_redact_annot(b["bbox"], fill=False)
        p.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                           graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                           text=pymupdf.PDF_REDACT_TEXT_REMOVE)


@tool("pdf_to_ppt")
def pdf_to_ppt(ctx):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    from ..fonts import LABELS, guess_family, style_from

    mode = ctx.opt("mode", "editable")
    EMU = 12700  # 1 pt
    paths = []
    for e in ctx.entries:
        doc = ctx.open_pdf(e)
        prs = Presentation()
        first = doc[0].rect
        prs.slide_width = Emu(int(first.width * EMU))
        prs.slide_height = Emu(int(first.height * EMU))
        blank = prs.slide_layouts[6]
        for i in range(doc.page_count):
            page = doc[i]
            slide = prs.slides.add_slide(blank)
            sx = first.width / page.rect.width
            sy = first.height / page.rect.height
            if mode == "image":
                pix = page.get_pixmap(dpi=200, alpha=False)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0,
                                         prs.slide_width, prs.slide_height)
                continue
            # Düzenlenebilir: metinsiz arka plan + üstte metin kutuları
            tmp = pymupdf.open()
            tmp.insert_pdf(doc, from_page=i, to_page=i)
            if tmp[0].rotation:
                tmp[0].remove_rotation()
            blocks = tmp[0].get_text("dict")["blocks"]
            _strip_text(tmp)
            pix = tmp[0].get_pixmap(dpi=170, alpha=False)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0,
                                     prs.slide_width, prs.slide_height)
            for b in blocks:
                if b["type"] != 0:
                    continue
                for line in b["lines"]:
                    spans = [s for s in line["spans"] if s["text"].strip()]
                    if not spans or abs(line["dir"][1]) > 0.01:
                        continue
                    x0, y0, x1, y1 = line["bbox"]
                    box = slide.shapes.add_textbox(Emu(int(x0 * sx * EMU)), Emu(int(y0 * sy * EMU)),
                                                   Emu(int((x1 - x0 + 4) * sx * EMU)),
                                                   Emu(int((y1 - y0) * sy * EMU)))
                    tf = box.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    para = tf.paragraphs[0]
                    for s in line["spans"]:
                        run = para.add_run()
                        run.text = s["text"]
                        f = run.font
                        f.size = Pt(max(1, s["size"] * sy))
                        bold, italic = style_from(s["font"], s["flags"])
                        f.bold, f.italic = bold, italic
                        f.name = LABELS[guess_family(s["font"], s["flags"])].split(" /")[0]
                        c = s["color"]
                        f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
        dst = ctx.out(stem(e.name) + ".pptx")
        prs.save(dst)
        paths.append(dst)
    ctx.result_name = "powerpoint.zip"
    return paths


_NUM_TR = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+(,\d+)?$")
_NUM_EN = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")


def _cell(v, numbers: str):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    t = s.replace("₺", "").replace("TL", "").replace("$", "").replace("€", "").replace("%", "").strip()
    try:
        if numbers == "tr" and _NUM_TR.match(t):
            return float(t.replace(".", "").replace(",", "."))
        if numbers == "en" and _NUM_EN.match(t):
            return float(t.replace(",", ""))
    except ValueError:
        pass
    return s


@tool("pdf_to_excel")
def pdf_to_excel(ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    numbers = ctx.opt("numbers", "tr")
    layout = ctx.opt("layout", "per_table")
    paths = []
    for e in ctx.entries:
        doc = ctx.open_pdf(e)
        wb = Workbook()
        wb.remove(wb.active)
        single = wb.create_sheet("Tablolar") if layout == "single" else None
        row_ptr = 1
        found = 0
        for i, page in enumerate(doc):
            try:
                tables = page.find_tables().tables
            except Exception:
                tables = []
            for k, t in enumerate(tables):
                rows = t.extract()
                if not rows:
                    continue
                found += 1
                ws = single or wb.create_sheet(f"S{i + 1}-T{k + 1}")
                start = row_ptr if single else 1
                if single:
                    ws.cell(start, 1, f"Sayfa {i + 1} – Tablo {k + 1}").font = Font(bold=True)
                    start += 1
                for r, row in enumerate(rows):
                    for c, v in enumerate(row):
                        cell = ws.cell(start + r, c + 1, _cell(v, numbers))
                        if r == 0:
                            cell.font = Font(bold=True)
                if single:
                    row_ptr = start + len(rows) + 1
        if not found:  # tablo yoksa: satırları metin sütunlarına böl
            ws = wb.create_sheet("Metin")
            r = 1
            for i, page in enumerate(doc):
                for line in page.get_text("text").splitlines():
                    if not line.strip():
                        continue
                    for c, v in enumerate(re.split(r"\s{2,}|\t", line.strip())):
                        ws.cell(r, c + 1, _cell(v, numbers))
                    r += 1
            ctx.extra["warning"] = "Tablo bulunamadı; metin satırlar halinde aktarıldı."
        for ws in wb.worksheets:
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(8, width + 2))
        dst = ctx.out(stem(e.name) + ".xlsx")
        wb.save(dst)
        paths.append(dst)
    ctx.result_name = "excel.zip"
    return paths


def _srgb_icc() -> bytes:
    from PIL import ImageCms
    return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


@tool("pdf_to_pdfa")
def pdf_to_pdfa(ctx):
    import pikepdf

    part = int(ctx.opt("part", "2"))
    paths, warnings_ = [], []
    for e in ctx.entries:
        src = ctx.pdf_path(e)
        dst = ctx.out(f"{stem(e.name)}_pdfa.pdf")
        if backends.capabilities()["ghostscript"]:
            backends.gs_pdfa(src, dst, part)
            ctx.extra["warning"] = "PDF/A dönüşümü yapıldı; arşiv uygunluğu bağımsız bir doğrulayıcıyla (ör. veraPDF) denetlenmelidir."
            paths.append(dst)
            continue
        # Ghostscript yoksa: en iyi çaba ile PDF/A işaretleme
        doc = pymupdf.open(src)
        missing = sorted({f[3] for p in doc for f in p.get_fonts() if f[1] in ("n/a", "")})
        if missing:
            warnings_.append(f"{e.name}: gömülü olmayan yazı tipleri: {', '.join(missing[:5])}")
        try:
            doc.scrub(javascript=True, attached_files=part != 3, embedded_files=part != 3,
                      metadata=False, xml_metadata=False, thumbnails=False, reset_fields=False,
                      reset_responses=False, clean_pages=False, hidden_text=False,
                      redactions=False, remove_links=False, redact_images=0)
        except Exception:
            pass
        tmp = ctx.out("pdfa_tmp.pdf")
        doc.save(tmp, garbage=3, deflate=True)
        doc.close()
        with pikepdf.open(tmp) as pdf:
            icc = pdf.make_stream(_srgb_icc())
            icc.N = 3
            intent = pikepdf.Dictionary(
                Type=pikepdf.Name.OutputIntent, S=pikepdf.Name.GTS_PDFA1,
                OutputConditionIdentifier=pikepdf.String("sRGB IEC61966-2.1"),
                Info=pikepdf.String("sRGB IEC61966-2.1"), DestOutputProfile=icc)
            pdf.Root.OutputIntents = pikepdf.Array([intent])
            if "/Info" not in pdf.trailer or not pdf.docinfo.get("/Title"):
                pdf.docinfo["/Title"] = stem(e.name)
            with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
                meta.load_from_docinfo(pdf.docinfo)
                meta["pdfaid:part"] = str(part)
                meta["pdfaid:conformance"] = "B"
            pdf.save(dst, min_version="1.7" if part > 1 else "1.4", fix_metadata_version=True)
        paths.append(dst)
    if warnings_:
        ctx.extra["warning"] = ("Ghostscript kurulu olmadığı için dönüşüm en iyi çaba ile yapıldı. "
                                + " ".join(warnings_))
    elif not backends.capabilities()["ghostscript"]:
        ctx.extra["warning"] = ("PDF/A işaretleri eklendi. Tam uyumluluk doğrulaması için "
                                "Ghostscript kurmanı öneririm.")
    ctx.result_name = "pdfa.zip"
    return paths


@tool("pdf_to_text")
def pdf_to_text(ctx):
    fmt = ctx.opt("format", "txt")
    paths = []
    for e in ctx.entries:
        doc = ctx.open_pdf(e)
        if fmt == "html":
            body = "\n".join(f'<section class="page" data-page="{i + 1}">{p.get_text("xhtml")}</section>'
                             for i, p in enumerate(doc))
            html = ("<!doctype html><html lang='tr'><head><meta charset='utf-8'>"
                    f"<title>{stem(e.name)}</title><style>body{{max-width:900px;margin:2rem auto;"
                    "font-family:system-ui;line-height:1.5}}.page{{border-bottom:1px solid #ccc;"
                    f"padding:1rem 0}}img{{max-width:100%}}</style></head><body>{body}</body></html>")
            p = ctx.out(stem(e.name) + ".html")
            p.write_text(html, encoding="utf-8")
        else:
            text = "\n\n".join(f"--- Sayfa {i + 1} ---\n{p.get_text('text', sort=True)}"
                               for i, p in enumerate(doc))
            p = ctx.out(stem(e.name) + ".txt")
            p.write_text(text, encoding="utf-8-sig")
        paths.append(p)
    ctx.result_name = "metin.zip"
    return paths

