"""Smoke-test actual Linux Office converters and Turkish fonts inside the image."""
from pathlib import Path
import tempfile
import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from app.backends import office_to_pdf
from app.fonts import get_font

with tempfile.TemporaryDirectory(prefix="container-smoke-") as folder:
    root = Path(folder)
    word = Document()
    word.add_paragraph("PDF Atolye Turkish: ğ ş ı İ ç ö ü")
    word.save(root / "word.docx")
    excel = Workbook()
    excel.active["A1"] = "PDF Atolye"
    excel.save(root / "excel.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = "PDF Atolye"
    slides.save(root / "slides.pptx")
    for filename, kind in (("word.docx", "word"), ("excel.xlsx", "excel"), ("slides.pptx", "powerpoint")):
        result = office_to_pdf(root / filename, kind, root / f"{kind}.pdf")
        with pymupdf.open(result) as pdf:
            assert pdf.page_count > 0, filename
            assert "PDF Atolye" in " ".join(p.get_text() for p in pdf), filename
        print(f"PASS {kind} -> PDF")
    font = get_font()
    assert all(font.has_glyph(ord(c)) for c in "ğşıİçöü")
    print("PASS Turkish font glyphs")
