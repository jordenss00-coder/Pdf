import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from dataclasses import replace
from unittest.mock import patch

os.environ.update(PDF_MODE="hosted", PDF_ALLOWED_HOSTS="localhost,127.0.0.1",
                  PDF_SESSION_SECRET="test-secret-only-not-for-deployment-123456789",
                  PDF_ACCESS_PASSWORD="test-password-only")

import pymupdf
from fastapi.testclient import TestClient
from app import access, jobs, main, store
from app.settings import settings, Settings
from app.tools import run
from app.util import UserError


def pdf_bytes():
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), "Example heading", fontsize=22)
        page.insert_text((72, 110), "First paragraph", fontsize=11)
        page.insert_text((72, 135), "Second paragraph", fontsize=11)
        doc.new_page().insert_text((72, 72), "Second page")
        return doc.tobytes()


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pdf-tests-")
        self.data_patch = patch.object(store, "DATA", Path(self.temp.name))
        self.data_patch.start()
        main._render_cache.clear()
        main._login_attempts.clear()
        self.client = TestClient(main.app, base_url="https://localhost")
        self.client.__enter__()
        self.assertEqual(self.client.post("/api/login", json={"password": "test-password-only"}).status_code, 200)

    def tearDown(self):
        self.client.__exit__(None, None, None)
        store._entries.clear()
        store._active.clear()
        self.data_patch.stop()
        self.temp.cleanup()  # Fails on Windows if document handles remain open.

    def upload(self, client=None, name="sample.pdf"):
        r = (client or self.client).post("/api/upload", files={"files": (name, pdf_bytes(), "application/pdf")})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()[0]["id"]

    def process(self, tool, ids, options=None):
        return self.client.post("/api/process/" + tool, json={"files": ids, "options": options or {}})

    def test_health_static_and_capabilities(self):
        for url in ("/", "/api/health", "/api/runtime", "/static/js/app.js", "/static/css/app.css", "/static/vendor/lucide.min.js", "/static/vendor/Sortable.min.js"):
            self.assertEqual(self.client.get(url).status_code, 200, url)
        caps = self.client.get("/api/capabilities").json()
        self.assertEqual(caps["mode"], "hosted")
        self.assertFalse(caps["ai"])
        self.assertFalse(caps["browser"])

    def test_updates_are_off_outside_desktop(self):
        self.assertFalse(self.client.get("/api/runtime").json()["updates"])
        self.assertFalse(self.client.get("/api/update?auto=true").json()["supported"])
        for path in ("/api/update/check", "/api/update/download", "/api/update/auto"):
            self.assertEqual(self.client.post(path, json={}).status_code, 404)

    def test_login_cookie_and_unauthorized(self):
        with TestClient(main.app, base_url="https://localhost") as other:
            self.assertEqual(other.get("/api/capabilities").status_code, 401)
            self.assertEqual(other.post("/api/login", json={"password": "wrong"}).status_code, 401)
            result = other.post("/api/login", json={"password": "test-password-only"})
            cookie = result.headers["set-cookie"].lower()
            for attribute in ("httponly", "secure", "samesite=strict"):
                self.assertIn(attribute, cookie)

    def test_cross_session_ownership_all_routes(self):
        fid = self.upload()
        with TestClient(main.app, base_url="https://localhost") as other:
            other.post("/api/login", json={"password": "test-password-only"})
            for suffix in ("", "/download", "/page/0", "/text/0", "/meta", "/fields"):
                self.assertEqual(other.get(f"/api/files/{fid}{suffix}").status_code, 404)
            for suffix, body in (("password", {"password": "abc"}), ("search", {}), ("detect-fields", {})):
                self.assertEqual(other.post(f"/api/files/{fid}/{suffix}", json=body).status_code, 404)
            self.assertEqual(other.delete(f"/api/files/{fid}").status_code, 404)
            self.assertEqual(other.post("/api/process/rotate", json={"files": [fid]}).status_code, 404)

    def test_round_trip_api_process_and_download(self):
        fid = self.upload()
        for suffix in ("/page/0", "/text/0", "/fields", "/meta", "/download"):
            response = self.client.get(f"/api/files/{fid}{suffix}")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["cache-control"], "no-store")
        r = self.process("rotate", [fid], {"angle": 90})
        self.assertEqual(r.status_code, 200, r.text)
        output = self.client.get(f"/api/files/{r.json()['result']['id']}/download")
        with pymupdf.open(stream=output.content) as doc:
            self.assertEqual(doc.page_count, 2)
            self.assertEqual(doc[0].rotation, 90)
        self.assertEqual(list((store.DATA / "_work").iterdir()), [])

    def test_host_origin_and_body_guards(self):
        self.assertEqual(self.client.get("/api/health", headers={"Host": "attacker.test"}).status_code, 403)
        self.assertEqual(self.client.post("/api/login", json={"password": "test-password-only"}, headers={"Origin": "https://attacker.test"}).status_code, 403)
        self.assertEqual(self.client.post("/api/login", content=b"x", headers={"Content-Length": "9000000"}).status_code, 413)

    def test_chunked_limit_without_content_length(self):
        with patch.object(access, "settings", replace(settings, max_upload_mb=1)):
            response = self.client.post("/api/upload", content=(b"x" * 1024 * 1024 for _ in range(3)), headers={"Content-Type": "application/octet-stream"})
        self.assertEqual(response.status_code, 413)

    def test_upload_count_and_size(self):
        with patch.object(main, "settings", replace(settings, max_files=1)):
            r = self.client.post("/api/upload", files=[("files", ("a.pdf", pdf_bytes())), ("files", ("b.pdf", pdf_bytes()))])
            self.assertEqual(r.status_code, 413)
        with patch.object(main, "settings", replace(settings, max_upload_mb=1)):
            r = self.client.post("/api/upload", files={"files": ("a.pdf", b"x" * (1024 * 1024 + 1))})
            self.assertEqual(r.status_code, 413)
        self.assertEqual(len(store._entries), 0)

    def test_page_limit(self):
        with patch.object(store, "settings", replace(settings, max_pages=1)):
            response = self.client.post("/api/upload", files={"files": ("a.pdf", pdf_bytes())})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(store._entries)

    def test_settings_html_ai_and_certificate_guards(self):
        self.assertEqual(self.client.post("/api/settings", json={"anthropic_api_key": "x"}).status_code, 403)
        self.assertEqual(self.process("html_to_pdf", [], {"url": "http://127.0.0.1"}).status_code, 403)
        self.assertEqual(self.process("ai_summary", []).status_code, 403)
        fid = self.upload()
        self.assertEqual(self.process("sign", [fid], {"certificate": {"path": "/etc/passwd"}}).status_code, 400)

    def test_delete_logout_and_restart_persistence(self):
        fid = self.upload()
        store.init()
        self.assertEqual(self.client.get(f"/api/files/{fid}").status_code, 200)
        response = self.client.post("/api/logout")
        self.assertEqual(response.status_code, 200)
        self.assertFalse((store.DATA / fid).exists())
        self.assertEqual(self.client.get(f"/api/files/{fid}").status_code, 401)

    def test_expiry_cleanup(self):
        fid = self.upload()
        store._entries[fid].created = time.time() - store.MAX_AGE_SECONDS - 1
        self.assertEqual(self.client.get(f"/api/files/{fid}").status_code, 404)
        store.cleanup()
        self.assertFalse((store.DATA / fid).exists())

    def test_failed_job_cleanup_and_busy(self):
        fid = self.upload()
        r = self.process("does_not_exist", [fid])
        self.assertEqual(r.status_code, 400)
        self.assertEqual(list((store.DATA / "_work").iterdir()), [])
        store._active.add(fid)
        self.assertEqual(self.client.delete(f"/api/files/{fid}").status_code, 400)
        self.assertEqual(self.process("rotate", [fid]).status_code, 409)

    def test_session_quota_rolls_back(self):
        fid = self.upload()
        with patch.object(store, "settings", replace(settings, session_mb=0)):
            r = self.client.post("/api/upload", files={"files": ("second.pdf", pdf_bytes())})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(list(store._entries), [fid])

    def test_core_tools_release_files_and_permissions(self):
        entry = store.add_bytes(pdf_bytes(), "direct.pdf")
        second = store.add_bytes(pdf_bytes(), "second.pdf")
        cases = [("merge", [entry, second], {}, 4), ("split", [entry], {"mode": "all"}, 1),
                 ("extract_pages", [entry], {"pages": "1"}, 1), ("remove_pages", [entry], {"pages": "2"}, 1),
                 ("compress", [entry], {}, 2), ("reverse", [entry], {}, 2),
                 ("page_numbers", [entry], {}, 2), ("flatten", [entry], {}, 2)]
        for tool, entries, options, count in cases:
            with self.subTest(tool=tool):
                paths, _ = run(tool, entries, options)
                for path in paths:
                    with pymupdf.open(path) as doc:
                        self.assertEqual(doc.page_count, count)
        paths, _ = run("protect", [entry], {"password": "secret", "permissions": []})
        with pymupdf.open(paths[0]) as doc:
            self.assertTrue(doc.needs_pass)
            self.assertTrue(doc.authenticate("secret"))
            self.assertFalse(doc.permissions & pymupdf.PDF_PERM_PRINT)
            self.assertFalse(doc.permissions & pymupdf.PDF_PERM_COPY)
        protected = store.add_path(paths[0], move=False)
        paths, _ = run("unlock", [protected], {"password": "secret"})
        with pymupdf.open(paths[0]) as doc:
            self.assertFalse(doc.needs_pass)
        store.delete(entry.id)  # Also verifies no open source handle on Windows.

    def test_markdown_and_empty_scan_message(self):
        entry = store.add_bytes(pdf_bytes(), "markdown.pdf")
        paths, _ = run("pdf_to_markdown", [entry], {})
        text = paths[0].read_text(encoding="utf-8")
        self.assertIn("# Example heading", text)
        self.assertIn("First paragraph", text)
        self.assertIn("Sayfa 2", text)
        with pymupdf.open() as blank:
            blank.new_page()
            entry = store.add_bytes(blank.tobytes(), "blank.pdf")
        _, ctx = run("pdf_to_markdown", [entry], {})
        self.assertIn("OCR", ctx.extra["warning"])

    def test_deadline_terminates_worker(self):
        folder = store.work_dir()
        with patch.object(jobs, "settings", replace(settings, job_timeout=0.001)):
            with self.assertRaisesRegex(UserError, "saniye"):
                jobs.execute("rotate", [], {}, folder)

    def test_redaction_removes_text_and_find_replace(self):
        entry = store.add_bytes(pdf_bytes(), "content.pdf")
        paths, _ = run("redact", [entry], {"terms": ["First paragraph"]})
        with pymupdf.open(paths[0]) as doc:
            text = "".join(p.get_text() for p in doc)
            self.assertNotIn("First paragraph", text)
            self.assertIn("Second paragraph", text)
            for i in range(1, doc.xref_length()):
                if doc.xref_is_stream(i):
                    self.assertNotIn(b"First paragraph", doc.xref_stream(i))
        paths, ctx = run("find_replace", [entry], {"find": "First", "replace": "New"})
        self.assertEqual(ctx.extra["replaced"], 1)
        with pymupdf.open(paths[0]) as doc:
            self.assertIn("New", doc[0].get_text())
            self.assertNotIn("First", doc[0].get_text())

    def test_forms_create_and_flatten(self):
        entry = store.add_bytes(pdf_bytes(), "form.pdf")
        paths, _ = run("form", [entry], {"new_fields": [{"name": "name", "type": "text", "page": 0,
            "rect": {"x": 72, "y": 200, "w": 160, "h": 30}, "value": "Example User"}]})
        with pymupdf.open(paths[0]) as doc:
            self.assertEqual(next(doc[0].widgets()).field_value, "Example User")
        form_entry = store.add_path(paths[0], move=False)
        paths, _ = run("form", [form_entry], {"flatten": True})
        with pymupdf.open(paths[0]) as doc:
            self.assertFalse(list(doc[0].widgets() or []))
            self.assertIn("Example User", " ".join(doc[0].get_text().split()))

    def test_office_exports_and_formula_text(self):
        from openpyxl import load_workbook
        from pptx import Presentation
        from docx import Document
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72, 72), "=1+1")
            entry = store.add_bytes(doc.tobytes(), "formula.pdf")
        paths, _ = run("pdf_to_excel", [entry], {})
        wb = load_workbook(paths[0], data_only=False)
        cell = wb["Metin"]["A1"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")
        wb.close()
        entry = store.add_bytes(pdf_bytes(), "exports.pdf")
        paths, _ = run("pdf_to_ppt", [entry], {"mode": "image"})
        self.assertEqual(len(Presentation(paths[0]).slides), 2)
        paths, _ = run("pdf_to_word", [entry], {})
        self.assertIn("First paragraph", "\n".join(p.text for p in Document(paths[0]).paragraphs))

    def test_ocr_rasterized_page(self):
        with pymupdf.open(stream=pdf_bytes()) as source:
            pixels = source[0].get_pixmap(dpi=150)
        with pymupdf.open() as scan:
            page = scan.new_page()
            page.insert_image(page.rect, pixmap=pixels)
            entry = store.add_bytes(scan.tobytes(), "scan.pdf")
        paths, ctx = run("ocr", [entry], {"language": "eng", "dpi": 150})
        self.assertGreater(ctx.extra["words"], 0)
        with pymupdf.open(paths[0]) as doc:
            self.assertIn("paragraph", doc[0].get_text().lower())

    def test_image_round_trip_and_compare(self):
        entry = store.add_bytes(pdf_bytes(), "images.pdf")
        paths, _ = run("pdf_to_images", [entry], {"format": "png", "dpi": 72, "pages": "1"})
        self.assertEqual(len(paths), 1)
        image = store.add_path(paths[0], move=False)
        paths, _ = run("images_to_pdf", [image], {})
        with pymupdf.open(paths[0]) as doc:
            self.assertEqual(doc.page_count, 1)
            self.assertTrue(doc[0].get_images())
        _, ctx = run("compare", [entry, entry], {})
        self.assertEqual(ctx.extra["compare"]["inserted"], 0)
        self.assertEqual(ctx.extra["compare"]["deleted"], 0)

    def test_secret_validation_and_tampered_cookie(self):
        with self.assertRaises(RuntimeError):
            replace(settings, secret="short").validate()
        session = access.issue_session()
        self.assertIsNotNone(access.session_owner(session))
        self.assertIsNone(access.session_owner(session + "x"))

    def test_desktop_requires_private_token(self):
        desktop_settings = replace(settings, mode="local", desktop_token="desktop-test-token")
        with patch.object(access, "settings", desktop_settings), patch.object(main, "settings", desktop_settings):
            self.assertEqual(self.client.get("/api/runtime").status_code, 403)
            self.assertEqual(self.client.get("/desktop/wrong").status_code, 404)
            response = self.client.get("/desktop/desktop-test-token")
            self.assertEqual(response.status_code, 200)
            runtime = self.client.get("/api/runtime")
            self.assertEqual(runtime.status_code, 200)
            self.assertTrue(runtime.json()["desktop"])
            self.assertFalse(runtime.json()["login_required"])
            self.assertEqual(self.client.get("/api/health", headers={"X-PDF-Desktop": "wrong"}).status_code, 403)


if __name__ == "__main__":
    unittest.main()
